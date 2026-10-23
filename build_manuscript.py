"""Generate an editable research DRAFT directly from current experiment CSVs."""
from pathlib import Path
import pandas as pd
from docx import Document
from docx.shared import Inches,Pt
P=Path(__file__).resolve().parent; O=P/"outputs"
df=pd.read_csv(O/"run_level_metrics.csv")
ab=pd.read_csv(O/"ablation_runs.csv")
d=Document(); s=d.sections[0]
s.top_margin=s.bottom_margin=Inches(.75)
s.left_margin=s.right_margin=Inches(.9)
normal=d.styles["Normal"];normal.font.name="Times New Roman";normal.font.size=Pt(10)
normal.paragraph_format.line_spacing=2
for style in ["Heading 1","Heading 2"]:
 d.styles[style].font.name="Times New Roman";d.styles[style].font.size=Pt(11)
def para(s):d.add_paragraph(s)
def section(t):d.add_heading(t,1)
def table(headers,rows):
 t=d.add_table(rows=1,cols=len(headers))
 for i,h in enumerate(headers):t.rows[0].cells[i].text=h
 for record in rows:
  c=t.add_row().cells
  for i,v in enumerate(record):c[i].text=str(v)
 for row in t.rows:
  for cell in row.cells:
   for p in cell.paragraphs:p.paragraph_format.line_spacing=1

idlog=df[(df.model=="Logistic")&(df.split=="ID")].auc.mean()
idkd=df[(df.model=="K-DART")&(df.split=="ID")].auc.mean()
from scipy.stats import wilcoxon
pv=df[df.split=="ID"].pivot(index="seed",columns="model",values="auc")
pvalue=wilcoxon(pv["K-DART"]-pv["Logistic"]).pvalue

d.add_heading("K-DART: Knowledge-Guided AI for Data-Driven Digital Transformation Readiness and Enterprise Innovation",0)
para("Short Communication — SYNTHETIC RESEARCH DRAFT; NOT FOR JOURNAL SUBMISSION")
para("Author information, affiliations and corresponding email: [AUTHOR TO CONFIRM]")
section("Abstract")
para(f"This research-development study asks whether plausible rule-derived organizational features improve digital transformation readiness prediction. Twenty independently seeded synthetic enterprise datasets (2,400 observations each) were generated to compare logistic regression, gradient boosting and a knowledge-feature-augmented logistic model, K-DART. A separate shifted simulation assessed sensitivity to distribution changes. Mean within-distribution AUROC was {idkd:.4f} for K-DART and {idlog:.4f} for logistic regression (paired Wilcoxon p={pvalue:.4f}). The rule augmentation did not improve discrimination. Two experimental figures and two tables document baseline comparisons, ablation, and calibration. Results establish neither real-world readiness prediction nor validated decision utility; they illustrate the need to validate domain rules and enterprise outcomes independently before claims of knowledge-based advantage.")
para("Keywords: knowledge-based systems; digital transformation; enterprise innovation; data analytics; decision intelligence; AI")
section("1. Introduction")
para("Organizations considering digital transformation face heterogeneous technological, human and financial constraints. Machine-learning readiness predictions are not necessarily actionable: decision makers also need credible constraints and intervention priorities. The research question in this methodological pilot is whether adding rule-derived feature interactions improves predictive performance over conventional statistical and ensemble models. The current rules are hypothetical and should not be mistaken for independently elicited business knowledge. No priority or novelty claim is made without comprehensive prior-art verification.")
section("2. Methods")
para("The benchmark simulates six observed correlated enterprise indicators: infrastructure, workforce skills, innovation investment, leadership, process integration, and financial slack. Outcome labels are stochastically generated from three underlying latent factors with interaction terms. Each of 20 seeds produces a 70/30 stratified split of 2,400 records; each also generates 750 independently simulated shifted-distribution records. This is synthetic methodological testing, not survey-based enterprise validation.")
para("Standardized logistic regression and histogram gradient boosting serve as baselines. K-DART adds four prospective threshold or interaction features representing infrastructure insufficiency, the joint occurrence of infrastructure and leadership, infrastructure-process complementarity, and low workforce skill. All scaling is fit only on training data. The K-DART probability is σ(β₀+βᵀx+γᵀr(x)), where x comprises six observed variables and r(x) denotes four candidate knowledge-feature functions. These functions do not enforce hard constraints.")
para("AUROC, average precision, and Brier score assess predictive discrimination and calibration. A two-sided paired Wilcoxon test compares run-level K-DART and logistic AUROC. Threshold-only and interaction-only ablations isolate rule families. Random seeds and split definitions are implemented in experiment.py.")
section("3. Results")
rows=[]
for split in ["ID","SHIFT"]:
 for method in ["Logistic","GradientBoosting","K-DART"]:
  sub=df[(df.split==split)&(df.model==method)]
  rows.append([split,method,*[f"{sub[k].mean():.4f} ± {sub[k].std():.4f}" for k in ["auc","ap","brier"]]])
para("Table 1. Synthetic predictive performance over 20 runs (mean ± between-run SD).")
table(["Set","Model","AUROC","Avg. precision","Brier"],rows)
rows=[]
for var in ["All rules","No rule features","Threshold only","Interactions only"]:
 sub=ab[ab.variant==var];rows.append([var,f"{sub.auc.mean():.4f}",f"{sub.brier.mean():.4f}"])
para("Table 2. Knowledge-feature ablation (20 independent synthetic runs).")
table(["Variant","Mean AUROC","Mean Brier"],rows)
para(f"Mean synthetic AUROC was {idkd:.4f} for K-DART versus {idlog:.4f} for the logistic baseline. The paired test did not show a significant difference (p={pvalue:.4f}). The observed performance does not support a predictive advantage from the proposed rules.")
for i,cap in [(1,"AUROC model comparison with 95% between-run normal-approximation intervals; not confidence bounds for a real population."),(2,"Calibration on one synthetic held-out run; not cross-dataset average calibration.")]:
 d.add_picture(str(O/f"Figure_{i}.png"),width=Inches(5.4))
 para(f"Figure {i}. {cap}")
section("4. Discussion")
para("These results demonstrate a common risk in knowledge-augmented modeling: plausible engineered rules may not supply incremental information beyond correlated base predictors. A scientifically consequential knowledge-based decision model would require externally elicited constraints and evaluation of decisions, such as the expected cost of selecting unsuitable digital investments. This study does not measure those outcomes. Synthetic shift robustness is not proof of cross-sector generalization.")
section("5. Limitations and conclusion")
para("Labels, enterprise profiles and rule thresholds are simulated. No organizational measurements, expert panel, prospective intervention outcomes, or external validity checks are present. Hence K-DART is not a deployable readiness instrument. Negative results are reported without manufacturing a performance benefit. Future research should acquire approved enterprise data, define an independent readiness target, validate rules with experts, quantify decision utility and compare against stronger knowledge-based baselines.")
section("Data and software availability")
para("Synthetic observations and results are reproducible using experiment.py; source code and generated outputs are hosted in the project repository. Public repository release is not a substitute for depositing an archived research artifact with a persistent identifier where required.")
section("Declarations")
para("Author contributions, funding and competing-interest details require author confirmation. Generative AI assisted manuscript structuring and coding; the author must critically verify all content, references and disclosures before submission.")
section("References")
para("[1] S.M. Lundberg, S.-I. Lee, A unified approach to interpreting model predictions, Advances in Neural Information Processing Systems 30 (2017).")
para("[2] World Bank Group, Enterprise Surveys data, https://www.enterprisesurveys.org/en/data (accessed 23 October 2026).")
para("[3] Eurostat, Community Innovation Survey microdata, https://ec.europa.eu/eurostat/web/microdata/collections-research/community-innovation-survey (accessed 23 October 2026).")
out=P/"KDART_Short_Communication_RESEARCH_DRAFT.docx";d.save(out)
print("Generated",out)
